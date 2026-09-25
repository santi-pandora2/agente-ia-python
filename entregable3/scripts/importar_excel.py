"""Convierte el Excel de la clase en una base SQLite reconstruible."""

from __future__ import annotations

import argparse
import sqlite3
import warnings
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def texto(valor: object) -> str | None:
    if valor is None:
        return None
    limpio = str(valor).strip()
    return limpio or None


def fecha_iso(valor: object) -> str | None:
    if isinstance(valor, (datetime, date)):
        return valor.date().isoformat() if isinstance(valor, datetime) else valor.isoformat()
    return texto(valor)


def filas_de_datos(hoja, columnas: int):
    """La fila 4 contiene encabezados; los datos empiezan en la fila 5."""
    for numero, fila in enumerate(
        hoja.iter_rows(min_row=5, max_col=columnas, values_only=True), start=5
    ):
        if fila[0] is not None:
            yield numero, fila


def convertir(origen: Path, destino: Path) -> dict[str, int]:
    if not origen.exists():
        raise FileNotFoundError(f"No se encontró el Excel: {origen}")

    destino.parent.mkdir(parents=True, exist_ok=True)
    temporal = destino.with_suffix(".temporal.db")
    temporal.unlink(missing_ok=True)

    # data_only lee los valores calculados guardados. openpyxl no ejecuta macros.
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        libro = load_workbook(origen, data_only=True, keep_vba=True)
    requeridas = {"Productos", "Proveedores", "Inventario"}
    faltantes = requeridas.difference(libro.sheetnames)
    if faltantes:
        raise ValueError(f"Faltan hojas requeridas: {', '.join(sorted(faltantes))}")

    conteos = {"proveedores": 0, "productos": 0, "inventario": 0}
    conexion = sqlite3.connect(temporal)
    try:
        conexion.execute("PRAGMA foreign_keys = ON")
        conexion.executescript(
            """
            CREATE TABLE proveedores (
                id TEXT PRIMARY KEY,
                razon_social TEXT NOT NULL,
                nit TEXT,
                categoria TEXT,
                contacto TEXT,
                telefono TEXT,
                correo TEXT,
                ciudad TEXT,
                estado TEXT
            );
            CREATE TABLE productos (
                codigo TEXT PRIMARY KEY,
                nombre TEXT NOT NULL,
                categoria TEXT NOT NULL,
                proveedor_id TEXT NOT NULL REFERENCES proveedores(id),
                unidad TEXT,
                precio_costo INTEGER NOT NULL CHECK (precio_costo >= 0),
                precio_venta INTEGER NOT NULL CHECK (precio_venta >= 0)
            );
            CREATE TABLE inventario (
                id INTEGER PRIMARY KEY,
                producto_codigo TEXT NOT NULL REFERENCES productos(codigo),
                ubicacion TEXT,
                lote TEXT NOT NULL,
                stock INTEGER NOT NULL CHECK (stock >= 0),
                stock_minimo INTEGER NOT NULL CHECK (stock_minimo >= 0),
                fecha_vencimiento TEXT,
                estado TEXT,
                ultima_actualizacion TEXT,
                observaciones TEXT,
                fila_origen INTEGER NOT NULL,
                UNIQUE (producto_codigo, lote)
            );
            CREATE INDEX idx_productos_nombre ON productos(nombre);
            CREATE INDEX idx_inventario_vencimiento ON inventario(fecha_vencimiento);
            """
        )

        for _, fila in filas_de_datos(libro["Proveedores"], 14):
            conexion.execute(
                "INSERT INTO proveedores VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    texto(fila[0]), texto(fila[1]), texto(fila[2]), texto(fila[3]),
                    texto(fila[4]), texto(fila[5]), texto(fila[7]), texto(fila[9]),
                    texto(fila[13]),
                ),
            )
            conteos["proveedores"] += 1

        for _, fila in filas_de_datos(libro["Productos"], 10):
            conexion.execute(
                "INSERT INTO productos VALUES (?, ?, ?, ?, ?, ?, ?)",
                (
                    texto(fila[0]), texto(fila[1]), texto(fila[2]), texto(fila[3]),
                    texto(fila[5]), int(fila[6] or 0), int(fila[7] or 0),
                ),
            )
            conteos["productos"] += 1

        for numero, fila in filas_de_datos(libro["Inventario"], 15):
            conexion.execute(
                """INSERT INTO inventario
                   (producto_codigo, ubicacion, lote, stock, stock_minimo,
                    fecha_vencimiento, estado, ultima_actualizacion,
                    observaciones, fila_origen)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    texto(fila[0]), texto(fila[4]), texto(fila[5]), int(fila[6] or 0),
                    int(fila[7] or 0), fecha_iso(fila[10]), texto(fila[12]),
                    fecha_iso(fila[13]), texto(fila[14]), numero,
                ),
            )
            conteos["inventario"] += 1

        errores_fk = conexion.execute("PRAGMA foreign_key_check").fetchall()
        if errores_fk:
            raise ValueError(f"Se encontraron relaciones inválidas: {errores_fk[:5]}")
        conteos["bajo_minimo"] = conexion.execute(
            "SELECT COUNT(*) FROM inventario WHERE stock < stock_minimo"
        ).fetchone()[0]
        conexion.commit()
    except Exception:
        conexion.rollback()
        raise
    finally:
        conexion.close()
        libro.close()

    temporal.replace(destino)
    return conteos


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("origen", type=Path, help="Ruta del archivo .xlsm")
    parser.add_argument("destino", type=Path, help="Ruta del archivo .db")
    args = parser.parse_args()

    conteos = convertir(args.origen, args.destino)
    print(f"Base creada: {args.destino}")
    print(f"Proveedores: {conteos['proveedores']} (esperado: 75)")
    print(f"Productos: {conteos['productos']} (esperado: 1250)")
    print(f"Inventario: {conteos['inventario']} (esperado: 827)")
    print(f"Bajo el mínimo: {conteos['bajo_minimo']} (esperado: 73)")
    print("Relaciones inválidas: 0")


if __name__ == "__main__":
    main()
